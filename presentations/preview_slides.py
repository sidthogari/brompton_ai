"""Render pptx slides to PNG using shape geometry (approximate visual QA)."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

OUT = Path(__file__).resolve().parent / "output" / "previews"
OUT.mkdir(parents=True, exist_ok=True)

# 1920x1080 from 13.333 x 7.5 in
SCALE = 1920 / 13.333

try:
    FONT_R = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", 18)
    FONT_B = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf", 22)
except Exception:
    FONT_R = FONT_B = ImageFont.load_default()


def emu_in(x):
    return x / 914400


def rgb_of(color):
    if color is None:
        return (250, 246, 240)
    try:
        return (color[0], color[1], color[2])
    except Exception:
        return (28, 25, 23)


def fill_rgb(shape):
    try:
        if shape.fill.type is None:
            return None
        c = shape.fill.fore_color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return None


def render(pptx_path, prefix, max_slides=4):
    prs = Presentation(str(pptx_path))
    W, H = int(emu_in(prs.slide_width) * SCALE), int(emu_in(prs.slide_height) * SCALE)
    written = []
    for i, slide in enumerate(prs.slides, 1):
        if max_slides and i > max_slides:
            continue
        img = Image.new("RGB", (W, H), (250, 246, 240))
        draw = ImageDraw.Draw(img)
        for sh in slide.shapes:
            l = int(emu_in(sh.left) * SCALE)
            t = int(emu_in(sh.top) * SCALE)
            w = int(emu_in(sh.width) * SCALE)
            h = int(emu_in(sh.height) * SCALE)
            fill = fill_rgb(sh)
            if fill and sh.shape_type != MSO_SHAPE_TYPE.TABLE:
                draw.rectangle([l, t, l + w, t + h], fill=fill)
            if sh.has_text_frame:
                text = sh.text_frame.text.strip()
                if text:
                    draw.text((l + 6, t + 4), text[:180], fill=(28, 25, 23), font=FONT_R)
        dest = OUT / f"{prefix}_slide{i:02d}.png"
        img.save(dest)
        written.append(dest)
    return written


if __name__ == "__main__":
    base = Path(__file__).resolve().parent / "output"
    render(base / "Noli_Section1_Attribution_Architecture.pptx", "s1", max_slides=3)
    render(base / "Noli_Section2_SelfService_Revenue_Platform.pptx", "s2", max_slides=3)
    print("previews in", OUT)
