"""Shared visual system for the Noli Data Engineer interview decks."""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Noli-inspired editorial palette: warm paper, ink, clay, sage
INK = RGBColor(0x1C, 0x19, 0x17)
INK_SOFT = RGBColor(0x3A, 0x34, 0x30)
CREAM = RGBColor(0xFA, 0xF6, 0xF0)
PAPER = RGBColor(0xF3, 0xEC, 0xE3)
CARD = RGBColor(0xFF, 0xFC, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLAY = RGBColor(0xB0, 0x5C, 0x45)
CLAY_DEEP = RGBColor(0x8C, 0x43, 0x30)
SAGE = RGBColor(0x5F, 0x75, 0x62)
SAGE_DEEP = RGBColor(0x3E, 0x54, 0x45)
GOLD = RGBColor(0xB8, 0x95, 0x6A)
BLUSH = RGBColor(0xD4, 0xB8, 0xA8)
MUTED = RGBColor(0x7A, 0x71, 0x68)
LINE = RGBColor(0xE4, 0xD8, 0xC8)
DANGER = RGBColor(0xA3, 0x45, 0x3A)
OK = RGBColor(0x3E, 0x6B, 0x52)
NAVY = RGBColor(0x2C, 0x3A, 0x47)

FONT = "Calibri"
FONT_LIGHT = "Calibri Light"


def _set_run(run, size, bold=False, color=INK, font=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    # Help PowerPoint pick a Latin typeface consistently
    ea = rpr.find(qn("a:latin"))
    if ea is None:
        latin = parse_xml(
            f'<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
        )
        rpr.append(latin)


def add_textbox(slide, l, t, w, h, text, size=16, bold=False, color=INK, font=FONT,
                align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size, bold, color, font, italic)
    return box


def add_para(tf, text, size=14, bold=False, color=INK, font=FONT, space_before=0,
             space_after=4, align=PP_ALIGN.LEFT, italic=False, level=0):
    # First paragraph may already exist empty
    if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run(run, size, bold, color, font, italic)
    return p


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def stroke_shape(shape, color, pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def rect(slide, l, t, w, h, fill, line=None, line_pt=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, fill)
    if line:
        stroke_shape(s, line, line_pt)
    else:
        s.line.fill.background()
    return s


def rrect(slide, l, t, w, h, fill, line=None, line_pt=1.0, adj=0.08):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_shape(s, fill)
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    if line:
        stroke_shape(s, line, line_pt)
    else:
        s.line.fill.background()
    return s


def set_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def blank_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = CREAM if not dark else INK
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, bg)
    return slide


def content_chrome(slide, kicker, page, total, section="Section 1"):
    """Top kicker + bottom footer."""
    rect(slide, 0, 0, Inches(0.12), SLIDE_H, CLAY)
    add_textbox(slide, Inches(0.45), Inches(0.18), Inches(10), Inches(0.28),
                kicker.upper(), size=11, bold=True, color=CLAY, font=FONT)
    rect(slide, Inches(0.45), Inches(7.18), Inches(12.4), Inches(0.015), LINE)
    add_textbox(slide, Inches(0.45), Inches(7.20), Inches(8), Inches(0.24),
                f"Noli  ·  Data Engineer business case  ·  {section}",
                size=10, color=MUTED, font=FONT)
    add_textbox(slide, Inches(10.6), Inches(7.20), Inches(2.3), Inches(0.24),
                f"{page}  /  {total}", size=10, color=MUTED, font=FONT,
                align=PP_ALIGN.RIGHT)


def title_block(slide, title, subtitle=None, top=0.48):
    add_textbox(slide, Inches(0.45), Inches(top), Inches(12.4), Inches(0.7),
                title, size=28, bold=True, color=INK, font=FONT)
    if subtitle:
        add_textbox(slide, Inches(0.45), Inches(top + 0.58), Inches(12.4), Inches(0.4),
                    subtitle, size=14, color=INK_SOFT, font=FONT_LIGHT)


def card(slide, l, t, w, h, title, body, accent=CLAY, title_size=14, body_size=12):
    rrect(slide, l, t, w, h, CARD, LINE, 1.0, adj=0.06)
    rect(slide, l, t, Inches(0.08), h, accent)
    add_textbox(slide, l + Inches(0.22), t + Inches(0.12), w - Inches(0.34), Inches(0.36),
                title, size=title_size, bold=True, color=INK)
    box = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.46), w - Inches(0.34), h - Inches(0.56))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    _set_run(run, body_size, False, INK_SOFT, FONT)
    return box


def pill(slide, l, t, w, h, text, fill=CLAY, color=WHITE, size=11):
    rrect(slide, l, t, w, h, fill, adj=0.5)
    add_textbox(slide, l, t, w, h, text, size=size, bold=True, color=color,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def style_table(table, header_fill=INK, header_color=WHITE, zebra=PAPER):
    def _margins(cell):
        try:
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
        except Exception:
            pass

    for cell in table.rows[0].cells:
        _margins(cell)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for run in p.runs:
                _set_run(run, 11, True, header_color, FONT)
            cell.text_frame.word_wrap = True
            try:
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            except Exception:
                pass
    for i in range(1, len(table.rows)):
        row = table.rows[i]
        bg = CARD if i % 2 else zebra
        for cell in row.cells:
            _margins(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _set_run(run, 11, False, INK_SOFT, FONT)
                cell.text_frame.word_wrap = True
                try:
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                except Exception:
                    pass


def set_cell(table, r, c, text, bold=False):
    cell = table.cell(r, c)
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run(run, 11, bold, INK if r == 0 else INK_SOFT, FONT)


def bullet_card(slide, l, t, w, h, title, bullets, accent=CLAY):
    rrect(slide, l, t, w, h, CARD, LINE, 1.0, adj=0.05)
    rect(slide, l, t, w, Inches(0.08), accent)
    add_textbox(slide, l + Inches(0.18), t + Inches(0.16), w - Inches(0.3), Inches(0.34),
                title, size=14, bold=True, color=INK)
    box = slide.shapes.add_textbox(l + Inches(0.18), t + Inches(0.5), w - Inches(0.32), h - Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = "•  " + b
        _set_run(run, 12, False, INK_SOFT, FONT)
    return box
